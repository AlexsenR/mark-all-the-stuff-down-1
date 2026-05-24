import sys
from pathlib import Path

import docx

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
import convert as cv  # noqa: E402

CORPUS = REPO / "tests" / "corpus"

PNG_1X1 = bytes.fromhex(
    "89504E470D0A1A0A0000000D4948445200000001000000010804000000B51C0C02"
    "0000000B4944415478DA63FCFF1F0003030200EFBFA7DB0000000049454E44AE426082"
)


def test_docx_direct_renders_heading_and_paragraph(tmp_path):
    src = CORPUS / "simple_letter.docx"
    body = cv._convert_docx_direct(src)
    assert "Letter to Counterparty" in body
    assert "Dear Sir" in body


def test_docx_direct_renders_table_as_pipe_table(tmp_path):
    src = CORPUS / "Договор поставки № 47.docx"
    body = cv._convert_docx_direct(src)
    assert "Бумага А4" in body
    assert "Картриджи" in body
    assert "| Товар |" in body or "| Бумага А4 |" in body


def test_docx_direct_escapes_pipe_in_table_cells(tmp_path):
    src = tmp_path / "pipe.docx"
    d = docx.Document()
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Clause | Risk"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A | B"
    table.cell(1, 1).text = "C"
    d.save(src)

    body = cv._convert_docx_direct(src)

    assert "Clause \\| Risk" in body
    assert "A \\| B" in body


def test_dispatch_docx_writes_md_with_marker(run_convert, tmp_copy):
    src = tmp_copy(CORPUS / "simple_letter.docx")
    r = run_convert(src, env={"MTSD_FORCE_PREINSTALLED": "1"})
    assert r.returncode == 0, r.stderr
    md = src.with_suffix(".md")
    assert md.is_file()
    first = md.read_text(encoding="utf-8").splitlines()[0]
    assert first.startswith(f"<!-- mark-all-the-stuff-down v{cv.__version__}")
    assert "backend: python-docx" in first
    assert str(md) in r.stdout


def test_dispatch_docx_cyrillic_path(run_convert, tmp_copy):
    src = tmp_copy(CORPUS / "Договор поставки № 47.docx")
    r = run_convert(src, env={"MTSD_FORCE_PREINSTALLED": "1"})
    assert r.returncode == 0, r.stderr
    md_text = (src.with_suffix(".md")).read_text(encoding="utf-8")
    assert "Бумага А4" in md_text


def test_dispatch_docx_corrupt_errors_as_corrupted(run_convert, tmp_path):
    src = tmp_path / "corrupt.docx"
    src.write_bytes(b"not a valid docx package")
    r = run_convert(src, env={"MTSD_FORCE_PREINSTALLED": "1"})
    assert r.returncode == 1
    assert "ERROR corrupted" in r.stderr


def test_dispatch_docx_images_only_writes_placeholders(run_convert, tmp_path):
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    src = tmp_path / "images_only.docx"
    d = docx.Document()
    d.add_picture(str(image))
    d.save(src)

    r = run_convert(src, env={"MTSD_FORCE_PREINSTALLED": "1"})
    assert r.returncode == 0, r.stderr
    md_text = src.with_suffix(".md").read_text(encoding="utf-8")
    assert "<!-- embedded image: index=1 -->" in md_text
