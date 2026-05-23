"""Generate the full test corpus from code. Idempotent."""
from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path

import docx  # python-docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(
    os.environ.get("MTSD_CORPUS_ROOT", Path(__file__).resolve().parent / "corpus")
)
EXP = ROOT / "expected"


def _ensure_dirs():
    ROOT.mkdir(parents=True, exist_ok=True)
    EXP.mkdir(parents=True, exist_ok=True)


def _expected(name: str, **kw):
    payload = {
        "exit_code": kw.get("exit_code", 0),
        "error_type": kw.get("error_type"),
        "min_markdown_length": kw.get("min_markdown_length", 0),
        "required_substrings": kw.get("required_substrings", []),
        "forbidden_substrings": kw.get("forbidden_substrings", []),
        "backend": kw.get("backend"),
        "requires_ocr": kw.get("requires_ocr", False),
    }
    (EXP / f"{name}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


# ---------------- DOCX ----------------


def _docx_simple_letter():
    d = docx.Document()
    d.add_heading("Letter to Counterparty", level=1)
    d.add_paragraph("Dear Sir or Madam,")
    d.add_paragraph("Please find attached the draft agreement.")
    d.save(ROOT / "simple_letter.docx")
    _expected(
        "simple_letter.docx",
        min_markdown_length=40,
        required_substrings=["Letter to Counterparty", "Dear Sir"],
    )


def _docx_postavki():
    d = docx.Document()
    d.add_heading("Договор поставки № 47", level=1)
    d.add_paragraph("г. Москва, 25 мая 2026 года")
    t = d.add_table(rows=3, cols=3)
    t.rows[0].cells[0].text = "№"
    t.rows[0].cells[1].text = "Товар"
    t.rows[0].cells[2].text = "Цена"
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].text = "Бумага А4"
    t.rows[1].cells[2].text = "350"
    t.rows[2].cells[0].text = "2"
    t.rows[2].cells[1].text = "Картриджи"
    t.rows[2].cells[2].text = "4500"
    d.save(ROOT / "Договор поставки № 47.docx")
    _expected(
        "Договор поставки № 47.docx",
        min_markdown_length=100,
        required_substrings=["Договор поставки", "Бумага А4", "Картриджи"],
    )


def _docx_revisions():
    d = docx.Document()
    d.add_heading("Соглашение", level=1)
    d.add_paragraph("Сторона А и Сторона Б договорились о нижеследующем.")
    d.add_paragraph("Пункт 1. Предмет соглашения.")
    d.add_paragraph("Пункт 2. Срок действия — один год.")
    d.save(ROOT / "Соглашение (с правками).docx")
    _expected(
        "Соглашение (с правками).docx",
        min_markdown_length=60,
        required_substrings=["Соглашение", "Сторона А", "Пункт 1"],
    )


def _docx_dated():
    d = docx.Document()
    d.add_heading("Договор №1 от 25.05.2026", level=1)
    d.add_paragraph("Заключён между ООО «Альфа» и ИП Иванов.")
    d.save(ROOT / "Договор №1 от 25.05.2026.docx")
    _expected(
        "Договор №1 от 25.05.2026.docx",
        min_markdown_length=40,
        required_substrings=["Договор №1", "ООО «Альфа»"],
    )


# ---------------- DOC (legacy, unsupported) ----------------


def _doc_legacy():
    # We don't need a valid .doc - convert.py rejects by extension before opening.
    (ROOT / "report_1998.doc").write_bytes(
        b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\x00" * 64
    )
    _expected("report_1998.doc", exit_code=1, error_type="unsupported")


# ---------------- PDF ----------------


def _pdf_simple_memo():
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError:
        raise RuntimeError("reportlab required: pip install reportlab")

    font = "Helvetica"
    for candidate in (
        "/Library/Fonts/Arial Unicode.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ):
        if Path(candidate).is_file():
            pdfmetrics.registerFont(TTFont("Cyr", candidate))
            font = "Cyr"
            break
    c = canvas.Canvas(str(ROOT / "Меморандум.pdf"), pagesize=A4)
    c.setFont(font, 12)
    c.drawString(72, 800, "Меморандум")
    c.drawString(72, 780, "Дата: 25 мая 2026 года")
    c.drawString(72, 760, "Тема: Анализ договора поставки")
    c.drawString(72, 740, "Настоящий меморандум подготовлен по запросу клиента.")
    c.save()
    _expected(
        "Меморандум.pdf",
        min_markdown_length=30,
        required_substrings=["Меморандум"],
    )


def _pdf_multicolumn():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(ROOT / "статья_журнал.pdf"), pagesize=A4)
    for i, line in enumerate(
        [
            "Column A line 1",
            "Column A line 2",
            "Column A line 3",
            "Column A line 4",
            "Column A line 5",
            "Column A line 6",
        ]
    ):
        c.drawString(50, 800 - i * 16, line)
    for i, line in enumerate(
        [
            "Column B line 1",
            "Column B line 2",
            "Column B line 3",
            "Column B line 4",
            "Column B line 5",
            "Column B line 6",
        ]
    ):
        c.drawString(320, 800 - i * 16, line)
    c.save()
    _expected(
        "статья_журнал.pdf",
        min_markdown_length=80,
        required_substrings=["Column A", "Column B"],
    )


def _pdf_tables():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table

    c = canvas.Canvas(str(ROOT / "Бюджет проекта 2026.pdf"), pagesize=A4)
    data = [
        ["Item", "Q1", "Q2", "Q3", "Q4"],
        ["Salaries", "100", "100", "110", "120"],
        ["Office", "20", "20", "20", "25"],
        ["Travel", "5", "15", "10", "8"],
    ]
    t = Table(data)
    t.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])
    t.wrapOn(c, 400, 600)
    t.drawOn(c, 50, 600)
    t2 = Table(data)
    t2.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])
    t2.wrapOn(c, 400, 600)
    t2.drawOn(c, 50, 400)
    c.save()
    _expected(
        "Бюджет проекта 2026.pdf",
        min_markdown_length=80,
        required_substrings=["Salaries", "Travel"],
    )


def _pdf_scanned():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(
        str(ROOT / "Контракт от 25.05 — окончательная редакция.pdf"), pagesize=A4
    )
    c.setFillColorRGB(0.9, 0.9, 0.9)
    c.rect(50, 50, 500, 700, fill=1, stroke=0)
    c.showPage()
    c.save()
    _expected(
        "Контракт от 25.05 — окончательная редакция.pdf",
        exit_code=0,
        backend="tesseract-ocr",
        requires_ocr=True,
    )


def _pdf_protected():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    p = ROOT / "protected.pdf"
    c = canvas.Canvas(str(p), pagesize=A4, encrypt="ownerpw")
    c.drawString(72, 800, "Secret")
    c.save()
    _expected("protected.pdf", exit_code=1, error_type="protected")


# ---------------- XLSX ----------------


def _xlsx_clients():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    ws.append(["ID", "ФИО", "Договор", "Сумма"])
    ws.append([1, "Иванов И.И.", "Д-001", 100000])
    ws.append([2, "Петров П.П.", "Д-002", 250000])
    wb.save(ROOT / "Реестр клиентов 2026.xlsx")
    _expected(
        "Реестр клиентов 2026.xlsx",
        min_markdown_length=80,
        required_substrings=["## Клиенты", "Иванов И.И.", "250000"],
    )


def _xlsx_merged():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Финмодель"
    ws["A1"] = "Заголовок таблицы"
    ws.merge_cells("A1:D1")
    ws.append(["Период", "Выручка", "Расходы", "Прибыль"])
    ws.append([2026, 1000, 700, 300])
    wb.save(ROOT / "Финмодель v3.xlsx")
    _expected(
        "Финмодель v3.xlsx",
        min_markdown_length=60,
        required_substrings=["## Финмодель", "Заголовок таблицы", "Выручка"],
    )


def _xlsx_formulas_multisheet():
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "Расчёт"
    s1.append(["A", "B", "Сумма"])
    s1.append([10, 20, "=A2+B2"])
    s1["C2"].value = "=A2+B2"
    s2 = wb.create_sheet("Неустойка")
    s2.append(["День", "Ставка"])
    s2.append([1, 0.1])
    wb.save(ROOT / "Расчёт неустойки.xlsx")
    _expected(
        "Расчёт неустойки.xlsx",
        min_markdown_length=80,
        required_substrings=["## Расчёт", "## Неустойка", "=A2+B2"],
    )


# ---------------- XLS (legacy, unsupported in preinstalled bundle) ----------------


def _xls_legacy():
    try:
        import xlwt
    except ImportError as e:
        raise RuntimeError("xlwt required for .xls fixture: pip install xlwt") from e
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Legacy")
    ws.write(0, 0, "ID")
    ws.write(0, 1, "Название")
    ws.write(1, 0, 1)
    ws.write(1, 1, "Контрагент Альфа")
    ws.write(2, 0, 2)
    ws.write(2, 1, "Контрагент Бета")
    wb.save(str(ROOT / "legacy_export.xls"))
    _expected(
        "legacy_export.xls",
        min_markdown_length=20,
        required_substrings=["Контрагент Альфа"],
    )
    _expected("legacy_export.xls.preinstalled", exit_code=1, error_type="install-failed")


# ---------------- PPTX ----------------


def _pptx_corrected():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Презентация (исправлено)"
    s.placeholders[1].text = "Версия от 25 мая 2026 года"
    p.save(ROOT / "Презентация (исправлено).pptx")
    _expected(
        "Презентация (исправлено).pptx",
        min_markdown_length=20,
        required_substrings=["Презентация"],
    )


def _pptx_notes():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Кейс для клиента"
    body = s.placeholders[1].text_frame
    body.text = "Краткое изложение позиции"
    s.notes_slide.notes_text_frame.text = "Note: confidential."
    p.save(ROOT / "Кейс для клиента — заметки.pptx")
    _expected(
        "Кейс для клиента — заметки.pptx",
        min_markdown_length=20,
        required_substrings=["Кейс", "позиции"],
    )


def _pptx_with_images():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Презентация с картинками"
    png_path = ROOT / "_pixel.png"
    png_path.write_bytes(
        bytes.fromhex(
            "89504E470D0A1A0A0000000D49484452000000010000000108060000001F15C489"
            "0000000D49444154789C6300010000000500010D0A2DB40000000049454E44AE426082"
        )
    )
    s.shapes.add_picture(str(png_path), Inches(1), Inches(2), Inches(2), Inches(2))
    p.save(ROOT / "Презентация с картинками.pptx")
    png_path.unlink()
    _expected(
        "Презентация с картинками.pptx",
        min_markdown_length=20,
        required_substrings=["Презентация"],
        forbidden_substrings=[
            "This image shows",
            "The image depicts",
            "A photo of",
            "На изображении",
        ],
    )


# ---------------- RTF ----------------


def _rtf_postanovlenie():
    rtf = (
        r"{\rtf1\ansi\ansicpg1251\deff0"
        r"{\fonttbl{\f0\fnil\fcharset204 Arial;}}"
        r"\f0\fs24 "
        r"\'cf\'ee\'f1\'f2\'e0\'ed\'ee\'e2\'eb\'e5\'ed\'e8\'e5 \'b9 1234-\'f0"
        r"\par "
        r"\'ee\'f2 25 \'ec\'e0\'ff 2026 \'e3\'ee\'e4\'e0"
        r"}"
    )
    (ROOT / "Постановление № 1234-р.rtf").write_text(
        rtf, encoding="cp1252", errors="ignore"
    )
    _expected(
        "Постановление № 1234-р.rtf",
        min_markdown_length=10,
        required_substrings=["Постановление"],
    )
    _expected(
        "Постановление № 1234-р.rtf.preinstalled",
        exit_code=1,
        error_type="install-failed",
    )


# ---------------- HTML ----------------


def _html_court():
    html = """<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>Судебная практика</title></head><body>
<h1>Судебная практика по договорам поставки</h1>
<p>Решение от 10 мая 2026 года.</p>
<table><tr><th>Дело</th><th>Сторона</th></tr><tr><td>А40-1</td><td>ООО Альфа</td></tr></table>
</body></html>"""
    (ROOT / "судебная_практика.html").write_text(html, encoding="utf-8")
    _expected(
        "судебная_практика.html",
        min_markdown_length=40,
        required_substrings=["Судебная практика", "ООО Альфа"],
    )


def _html_report():
    html = """<!doctype html><html><body><h1>Report</h1>
<a href="https://example.com">link</a><p>Body text.</p></body></html>"""
    (ROOT / "report.html").write_text(html, encoding="utf-8")
    _expected(
        "report.html",
        min_markdown_length=10,
        required_substrings=["Report", "link"],
    )


# ---------------- EPUB ----------------


def _epub_codex():
    path = ROOT / "Гражданский_кодекс_РФ_часть_первая.epub"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED
        )
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package version="2.0" '
            'xmlns="http://www.idpf.org/2007/opf" unique-identifier="id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            "<dc:title>ГК РФ часть 1</dc:title>"
            '<dc:identifier id="id">urn:uuid:test</dc:identifier>'
            "<dc:language>ru</dc:language></metadata>"
            '<manifest><item id="ch1" href="ch1.xhtml" '
            'media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="ch1"/></spine></package>',
        )
        z.writestr(
            "OEBPS/ch1.xhtml",
            '<?xml version="1.0"?><!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml">'
            "<head><title>Гл. 1</title></head><body>"
            "<h1>Глава 1. Общие положения</h1>"
            "<p>Гражданское законодательство...</p></body></html>",
        )
    _expected(
        "Гражданский_кодекс_РФ_часть_первая.epub",
        min_markdown_length=20,
        required_substrings=["Глава 1"],
    )
    _expected(
        "Гражданский_кодекс_РФ_часть_первая.epub.preinstalled",
        exit_code=1,
        error_type="install-failed",
    )


def main():
    _ensure_dirs()
    _docx_simple_letter()
    _docx_postavki()
    _docx_revisions()
    _docx_dated()
    _doc_legacy()
    _pdf_simple_memo()
    _pdf_multicolumn()
    _pdf_tables()
    _pdf_scanned()
    _pdf_protected()
    _xlsx_clients()
    _xlsx_merged()
    _xlsx_formulas_multisheet()
    _xls_legacy()
    _pptx_corrected()
    _pptx_notes()
    _pptx_with_images()
    _rtf_postanovlenie()
    _html_court()
    _html_report()
    _epub_codex()
    print(f"Generated corpus in {ROOT}")


if __name__ == "__main__":
    main()
