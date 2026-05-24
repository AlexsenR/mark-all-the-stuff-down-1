#!/usr/bin/env python3
"""mark-all-the-stuff-down - convert office docs to Markdown (with OCR)."""
from __future__ import annotations

import argparse
import concurrent.futures
import contextlib
import hashlib
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path

__version__ = "0.1.0"
MARKER_PREFIX = "<!-- mark-all-the-stuff-down v"
MAX_BYTES = 50_000_000  # 50 MB (decimal: 1 KB = 1000 bytes)
PROBE_PAGES = 3
SCANNED_PDF_TEXT_THRESHOLD = 50  # chars/page average

SUPPORTED_EXTS = {
    ".bmp",
    ".docx",
    ".epub",
    ".html",
    ".jpeg",
    ".jpg",
    ".pdf",
    ".png",
    ".pptx",
    ".rtf",
    ".tif",
    ".tiff",
    ".xls",
    ".xlsx",
    ".webp",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}
UNSUPPORTED_DOC = {".doc"}

CHILD_ENV = {
    "LANG": "en_US.UTF-8",
    "LC_ALL": "en_US.UTF-8",
    "PYTHONIOENCODING": "utf-8",
}

MIT_SNIPPET = (
    "import sys; from markitdown import MarkItDown; "
    "sys.stdout.write(MarkItDown().convert(sys.argv[1]).text_content)"
)
PMU_SNIPPET = (
    "import pymupdf4llm,sys; "
    "sys.stdout.write(pymupdf4llm.to_markdown(sys.argv[1]))"
)


class Level(str, Enum):
    UVX = "uvx"
    PIP = "pip"
    PREINSTALLED = "preinstalled"


def _format_size(n: int) -> str:
    if n < 10 * 1000 * 1000:
        return f"{n / 1000:.1f} KB"
    return f"{n / 1_000_000:.1f} MB"


def _escape_md_table_cell(text: str) -> str:
    return text.replace("\n", " ").replace("|", "\\|")


def _cache_root() -> Path:
    if sys.platform == "win32":
        base = os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData" / "Local"))
        return Path(base) / "mtsd"
    return Path.home() / ".cache" / "mtsd"


def _has_our_marker(md_path: Path) -> bool:
    try:
        with md_path.open("r", encoding="utf-8") as fh:
            first = fh.readline()
    except OSError:
        return False
    return first.startswith(MARKER_PREFIX)


def _source_id(src: Path) -> str:
    return hashlib.sha256(str(src.resolve()).encode("utf-8")).hexdigest()[:16]


def _parse_marker_fields(line: str) -> dict[str, str]:
    if not line.startswith(MARKER_PREFIX) or not line.endswith("-->"):
        return {}
    inner = line[4:-3].strip()
    parts = [p.strip() for p in inner.split("|")]
    fields = {}
    for part in parts[1:]:
        if ": " in part:
            key, value = part.split(": ", 1)
            fields[key] = value
    return fields


def _marker_matches_source(md_path: Path, src: Path) -> bool:
    try:
        with md_path.open("r", encoding="utf-8") as fh:
            first = fh.readline().rstrip("\n")
    except OSError:
        return False
    return _parse_marker_fields(first).get("source_id") == _source_id(src)


def _output_candidate_for_source(candidate: Path, src: Path) -> bool:
    return not candidate.exists() or _marker_matches_source(candidate, src)


def compute_output_path(src: Path, reserved: set[Path] | None = None) -> Path:
    def available(candidate: Path) -> bool:
        if reserved is not None and candidate in reserved:
            return False
        return _output_candidate_for_source(candidate, src)

    sibling = src.with_suffix(".md")
    parent = src.parent
    if os.access(parent, os.W_OK):
        if available(sibling):
            return sibling
        mtsd = src.with_suffix(".mtsd.md")
        if available(mtsd):
            return mtsd
        ext_mtsd = src.with_name(f"{src.stem}{src.suffix}.mtsd.md")
        if available(ext_mtsd):
            return ext_mtsd
        counter = 2
        while True:
            numbered = src.with_name(f"{src.stem}{src.suffix}.{counter}.mtsd.md")
            if available(numbered):
                return numbered
            counter += 1
    cache = _cache_root()
    cache.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(str(src).encode("utf-8")).hexdigest()[:16]
    candidate = cache / f"{digest}.md"
    if reserved is None or candidate not in reserved:
        return candidate
    counter = 2
    while True:
        numbered = cache / f"{digest}.{counter}.md"
        if numbered not in reserved:
            return numbered
        counter += 1


def marker_line(src: Path, backend: str) -> str:
    stat = src.stat()
    mtime = (
        datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc)
        .astimezone()
        .replace(microsecond=0)
        .isoformat()
    )
    return (
        f"{MARKER_PREFIX}{__version__} | "
        f"source: {src.name} | source_id: {_source_id(src)} | "
        f"src_mtime: {stat.st_mtime_ns} | mtime: {mtime} | backend: {backend} -->"
    )


def _parse_marker_version(line: str) -> tuple[int, int, int] | None:
    if not line.startswith(MARKER_PREFIX):
        return None
    rest = line[len(MARKER_PREFIX) :]
    ver = rest.split(" ", 1)[0]
    try:
        a, b, c = (int(x) for x in ver.split("."))
        return a, b, c
    except ValueError:
        return None


def is_up_to_date(src: Path, md: Path) -> bool:
    if not md.is_file():
        return False
    try:
        with md.open("r", encoding="utf-8") as fh:
            first = fh.readline().rstrip("\n")
    except OSError:
        return False
    ver = _parse_marker_version(first)
    if ver is None:
        return False
    cur_major = int(__version__.split(".")[0])
    if ver[0] != cur_major:
        return False
    fields = _parse_marker_fields(first)
    if fields.get("source_id") != _source_id(src):
        return False
    try:
        if int(fields.get("src_mtime", "")) != src.stat().st_mtime_ns:
            return False
    except ValueError:
        return False
    return md.stat().st_mtime >= src.stat().st_mtime


def write_output(out: Path, marker: str, body: str) -> None:
    tmp = out.with_suffix(out.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8", newline="\n") as fh:
        fh.write(marker)
        if not marker.endswith("\n"):
            fh.write("\n")
        fh.write(body)
        if not body.endswith("\n"):
            fh.write("\n")
    tmp.replace(out)


def success_line(out: Path) -> str:
    size = out.stat().st_size
    return f"{out} ({_format_size(size)})"


def _path_hash(src: Path) -> str:
    return hashlib.sha256(str(src).encode("utf-8")).hexdigest()[:16]


def _log_path() -> Path:
    return _cache_root() / "log.jsonl"


def log_event(
    src: Path, *, backend: str, duration_ms: int, size_kb: int, exit_code: int
) -> None:
    if os.environ.get("MTSD_LOG") != "1":
        return
    rec = {
        "ts": datetime.now(tz=timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z"),
        "path_hash": _path_hash(src),
        "backend": backend,
        "duration_ms": duration_ms,
        "size_kb": size_kb,
        "exit": exit_code,
    }
    if os.environ.get("MTSD_LOG_VERBOSE") == "1":
        rec["path"] = str(src)
    log_path = _log_path()
    try:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, "a", encoding="utf-8") as fh:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    except OSError:
        pass


def _have_uv(no_uv: bool) -> bool:
    if no_uv:
        return False
    if os.environ.get("MTSD_FORCE_PREINSTALLED") == "1":
        return False
    return shutil.which("uv") is not None


def _have_network(host: str = "pypi.org", port: int = 443, timeout: float = 3.0) -> bool:
    if os.environ.get("MTSD_FORCE_PREINSTALLED") == "1":
        return False
    try:
        with socket.create_connection((host, port), timeout=timeout):
            return True
    except OSError:
        return False


def _run_child(
    cmd: list[str], *, timeout: int = 120, creationflags: int | None = None
) -> subprocess.CompletedProcess:
    env = {**os.environ, **CHILD_ENV}
    kwargs: dict = dict(
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=timeout,
        check=False,
    )
    if sys.platform == "win32":
        kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
    return subprocess.run(cmd, **kwargs)


def _argv_is_safe(p: Path) -> bool:
    try:
        str(p).encode(sys.getfilesystemencoding() or "utf-8")
        return True
    except UnicodeEncodeError:
        return False


@contextlib.contextmanager
def _with_safe_argv_path(src: Path):
    if _argv_is_safe(src):
        yield src
        return
    tmp = Path(tempfile.mkdtemp(prefix="mtsd_")) / f"input{src.suffix}"
    shutil.copy2(src, tmp)
    try:
        yield tmp
    finally:
        try:
            tmp.unlink()
            tmp.parent.rmdir()
        except OSError:
            pass


class BackendUnavailable(Exception):
    pass


class ConversionError(Exception):
    def __init__(self, kind: str, message: str):
        self.kind = kind
        self.message = message
        super().__init__(f"{kind}: {message}")


@dataclass
class ConversionResult:
    src: Path
    exit_code: int
    stdout_line: str | None = None
    error_kind: str | None = None
    error_message: str | None = None

    def stderr_line(self, *, include_path: bool) -> str | None:
        if self.error_kind is None:
            return None
        message = self.error_message or ""
        if include_path and self.error_kind != "not-found":
            message = f"{self.src}: {message}" if message else str(self.src)
        return f"ERROR {self.error_kind}: {message}"


_OUTPUT_LOCKS_GUARD = threading.Lock()
_OUTPUT_LOCKS: dict[Path, threading.Lock] = {}
_PIP_INSTALL_LOCK = threading.Lock()
_PIP_INSTALL_CACHE: dict[tuple[str, ...], bool] = {}


def _conversion_failure(
    src: Path, kind: str, message: str, *, exit_code: int = 1
) -> ConversionResult:
    return ConversionResult(
        src=src,
        exit_code=exit_code,
        error_kind=kind,
        error_message=message,
    )


def _output_lock_for(src: Path) -> threading.Lock:
    key = src.parent.resolve() / src.stem
    with _OUTPUT_LOCKS_GUARD:
        lock = _OUTPUT_LOCKS.get(key)
        if lock is None:
            lock = threading.Lock()
            _OUTPUT_LOCKS[key] = lock
        return lock


# ---------- DOCX ----------


def _convert_docx_direct(src: Path) -> str:
    try:
        import docx  # python-docx
        from docx.opc.exceptions import OpcError, PackageNotFoundError
        from zipfile import BadZipFile
    except ImportError as e:
        raise BackendUnavailable("python-docx not installed") from e
    try:
        from openpyxl.utils.exceptions import InvalidFileException

        docx_parse_errors = (PackageNotFoundError, OpcError, BadZipFile, InvalidFileException)
    except ImportError:
        docx_parse_errors = (PackageNotFoundError, OpcError, BadZipFile)
    try:
        d = docx.Document(str(src))
    except docx_parse_errors as e:
        raise ConversionError("corrupted", f"DOCX cannot be parsed: {e}") from e
    out_lines: list[str] = []
    body = d.element.body
    # Iterate top-level paragraphs and tables in document order.
    from docx.oxml.ns import qn

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = next((p for p in d.paragraphs if p._element is child), None)
            if para is None:
                continue
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading 1") or style == "title":
                out_lines.append(f"# {text}")
            elif style.startswith("heading 2"):
                out_lines.append(f"## {text}")
            elif style.startswith("heading 3"):
                out_lines.append(f"### {text}")
            else:
                out_lines.append(text)
            out_lines.append("")
        elif child.tag == qn("w:tbl"):
            tbl = next((t for t in d.tables if t._element is child), None)
            if tbl is None:
                continue
            rows = [
                [_escape_md_table_cell(c.text.strip()) for c in row.cells]
                for row in tbl.rows
            ]
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            rows = [r + [""] * (ncols - len(r)) for r in rows]
            out_lines.append("| " + " | ".join(rows[0]) + " |")
            out_lines.append("| " + " | ".join(["---"] * ncols) + " |")
            for r in rows[1:]:
                out_lines.append("| " + " | ".join(r) + " |")
            out_lines.append("")
    text = "\n".join(out_lines).rstrip()
    if not text:
        image_count = 0
        for rel in d.part.rels.values():
            target_part = getattr(rel, "target_part", None)
            content_type = getattr(target_part, "content_type", "")
            if content_type.startswith("image/"):
                image_count += 1
        if image_count:
            text = "\n".join(
                f"<!-- embedded image: index={i} -->"
                for i in range(1, image_count + 1)
            )
    return text + "\n"


# ---------- Generic markitdown invocation ----------


def _convert_via_markitdown_uvx(src: Path, extra: str) -> str:
    cmd = ["uvx", "--from", f"markitdown[{extra}]", "markitdown", str(src)]
    r = _run_child(cmd, timeout=180)
    if r.returncode != 0:
        raise BackendUnavailable(f"uvx markitdown[{extra}] failed: {r.stderr.strip()}")
    return r.stdout


def _convert_via_markitdown_pip(src: Path) -> str:
    cmd = [sys.executable, "-c", MIT_SNIPPET, str(src)]
    r = _run_child(cmd, timeout=180)
    if r.returncode != 0:
        raise BackendUnavailable(f"pip markitdown failed: {r.stderr.strip()}")
    return r.stdout


def _ensure_pip_installed(packages: list[str]) -> bool:
    key = tuple(packages)
    with _PIP_INSTALL_LOCK:
        cached = _PIP_INSTALL_CACHE.get(key)
        if cached is not None:
            return cached
        if not _have_network():
            _PIP_INSTALL_CACHE[key] = False
            return False
        cmd = [sys.executable, "-m", "pip", "install", "--user", "--quiet", *packages]
        r = _run_child(cmd, timeout=300)
        ok = r.returncode == 0
        _PIP_INSTALL_CACHE[key] = ok
        return ok


# ---------- PDF ----------


@dataclass
class PdfProbe:
    total_pages: int
    pages_probed: int
    table_count: int
    is_multicolumn: bool
    avg_chars_per_page: float
    is_scanned: bool
    complexity: str  # "simple" | "complex"


def _pdf_probe(src: Path) -> PdfProbe:
    try:
        import pdfplumber
        from pdfminer.pdfdocument import (
            PDFEncryptionError,
            PDFPasswordIncorrect,
            PDFTextExtractionNotAllowed,
        )
        from pdfminer.pdfexceptions import PDFException
        from pdfminer.pdfparser import PDFSyntaxError
        from pdfminer.psexceptions import PSException
        from pdfplumber.utils.exceptions import MalformedPDFException, PdfminerException
    except ImportError as e:
        raise BackendUnavailable("pdfplumber not installed") from e
    try:
        import pypdf
        from pypdf.errors import FileNotDecryptedError, PdfReadError
    except ImportError:
        pypdf = None
        PdfReadError = Exception  # type: ignore
        FileNotDecryptedError = Exception  # type: ignore

    if pypdf is not None:
        try:
            reader = pypdf.PdfReader(str(src))
            if reader.is_encrypted:
                try:
                    if reader.decrypt("") == 0:
                        raise ConversionError(
                            "protected",
                            "PDF is password-protected; ask sender for unprotected copy",
                        )
                except FileNotDecryptedError as e:
                    raise ConversionError(
                        "protected",
                        "PDF is password-protected; ask sender for unprotected copy",
                    ) from e
        except ConversionError:
            raise
        except (PdfReadError, ValueError, OSError) as e:
            raise ConversionError("corrupted", f"PDF cannot be parsed: {e}") from e

    try:
        with pdfplumber.open(str(src)) as pdf:
            total = len(pdf.pages)
            probe_n = min(PROBE_PAGES, total)
            pages = pdf.pages[:probe_n]
            table_count = 0
            total_chars = 0
            text_layer_chars = 0
            multicolumn = False
            for page in pages:
                try:
                    table_count += len(page.find_tables() or [])
                except Exception:
                    pass
                words = page.extract_words() or []
                total_chars += sum(len(w["text"]) for w in words)
                text_layer_chars += len((page.extract_text() or "").strip())
                multicolumn = multicolumn or _is_multicolumn(words, page.width)
            avg = total_chars / probe_n if probe_n else 0.0
            is_scanned = text_layer_chars == 0 and avg < SCANNED_PDF_TEXT_THRESHOLD
            complex_ = table_count >= 2 or multicolumn
            return PdfProbe(
                total,
                probe_n,
                table_count,
                multicolumn,
                avg,
                is_scanned,
                "complex" if complex_ else "simple",
            )
    except (PDFPasswordIncorrect, PDFEncryptionError, PDFTextExtractionNotAllowed) as e:
        raise ConversionError(
            "protected", "PDF is password-protected; ask sender for unprotected copy"
        ) from e
    except (
        PDFSyntaxError,
        PdfReadError,
        PDFException,
        PSException,
        PdfminerException,
        MalformedPDFException,
    ) as e:
        raise ConversionError("corrupted", f"PDF syntax error: {e}") from e
    except (OSError, ValueError, KeyError, AssertionError) as e:
        raise ConversionError("corrupted", f"PDF cannot be parsed: {e}") from e


def _is_multicolumn(words: list, page_width: float) -> bool:
    if not words or page_width <= 0:
        return False
    xs = sorted(w["x0"] for w in words)
    threshold = 0.30 * page_width
    cluster_starts = [xs[0]]
    cluster_counts = [1]
    for x in xs[1:]:
        if (
            x - cluster_starts[-1] - cluster_counts[-1] * 0 > threshold
            and (x - cluster_starts[-1]) >= threshold
        ):
            cluster_starts.append(x)
            cluster_counts.append(1)
        else:
            cluster_counts[-1] += 1
    big = [c for c in cluster_counts if c >= 5]
    return len(big) >= 2


def _convert_pdf_direct_simple(src: Path) -> str:
    try:
        import pdfplumber
        from pdfminer.pdfdocument import (
            PDFEncryptionError,
            PDFPasswordIncorrect,
            PDFTextExtractionNotAllowed,
        )
        from pdfminer.pdfexceptions import PDFException
        from pdfminer.pdfparser import PDFSyntaxError
        from pdfminer.psexceptions import PSException
        from pdfplumber.utils.exceptions import MalformedPDFException, PdfminerException
    except ImportError as e:
        raise BackendUnavailable("pdfplumber not installed") from e
    try:
        from pypdf.errors import PdfReadError
    except ImportError:
        PdfReadError = Exception  # type: ignore
    lines: list[str] = []
    try:
        with pdfplumber.open(str(src)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    lines.append(f"## Page {i}\n")
                    lines.append(text.strip())
                    lines.append("")
    except (PDFPasswordIncorrect, PDFEncryptionError, PDFTextExtractionNotAllowed) as e:
        raise ConversionError(
            "protected", "PDF is password-protected; ask sender for unprotected copy"
        ) from e
    except (
        PDFSyntaxError,
        PdfReadError,
        PDFException,
        PSException,
        PdfminerException,
        MalformedPDFException,
    ) as e:
        raise ConversionError("corrupted", f"PDF syntax error: {e}") from e
    except (OSError, ValueError, KeyError, AssertionError) as e:
        raise ConversionError("corrupted", f"PDF cannot be parsed: {e}") from e
    return "\n".join(lines).rstrip() + "\n"


# ---------- Dispatch ----------


def dispatch(src: Path, *, no_uv: bool) -> tuple[str, str]:
    """Return (markdown_body, backend_name)."""
    ext = src.suffix.lower()
    use_uv = _have_uv(no_uv)

    if ext == ".docx":
        return _do_docx(src, use_uv)
    if ext == ".pdf":
        return _do_pdf(src, use_uv)
    if ext == ".xlsx":
        return _do_xlsx(src, use_uv)
    if ext == ".pptx":
        return _do_pptx(src, use_uv)
    if ext == ".html":
        return _do_html(src, use_uv)
    if ext in {".xls", ".rtf", ".epub"}:
        return _do_markitdown_only(src, use_uv, ext.lstrip("."))
    if ext in IMAGE_EXTS:
        from ocr import OcrError, convert_image

        try:
            body, backend = convert_image(src)
        except OcrError as e:
            raise ConversionError(e.kind, e.message) from e
        return body, backend
    raise ConversionError("unsupported", f"format {ext} not in MVP scope")


def _do_docx(src: Path, use_uv: bool) -> tuple[str, str]:
    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, "docx"), "markitdown[docx]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed(["markitdown[docx]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), "markitdown[docx]/pip"
        except BackendUnavailable:
            pass
    try:
        return _convert_docx_direct(src), "python-docx"
    except BackendUnavailable as e:
        raise ConversionError(
            "install-failed", "docx requires python-docx or markitdown"
        ) from e


def _do_pdf(src: Path, use_uv: bool) -> tuple[str, str]:
    try:
        info = _pdf_probe(src)
    except BackendUnavailable as e:
        raise ConversionError("install-failed", "pdf requires pdfplumber") from e
    if info.is_scanned:
        from ocr import OcrError, convert_pdf_scanned

        try:
            body, backend = convert_pdf_scanned(src, info.total_pages)
        except OcrError as e:
            raise ConversionError(e.kind, e.message) from e
        return body, backend

    if info.complexity == "complex":
        if use_uv:
            with _with_safe_argv_path(src) as p:
                cmd = ["uvx", "--with", "pymupdf4llm", "python", "-c", PMU_SNIPPET, str(p)]
                r = _run_child(cmd, timeout=180)
                if r.returncode == 0:
                    return r.stdout, "pymupdf4llm/uvx"
        if _ensure_pip_installed(["pymupdf4llm"]):
            with _with_safe_argv_path(src) as p:
                r = _run_child([sys.executable, "-c", PMU_SNIPPET, str(p)], timeout=180)
                if r.returncode == 0:
                    return r.stdout, "pymupdf4llm/pip"
        body = _convert_pdf_direct_simple(src)
        print(
            "WARN pdf-complex-degraded: pymupdf4llm unavailable, "
            "falling back to pdfplumber",
            file=sys.stderr,
        )
        return body, "pdfplumber-fallback"

    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, "pdf"), "markitdown[pdf]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed(["markitdown[pdf]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), "markitdown[pdf]/pip"
        except BackendUnavailable:
            pass
    return _convert_pdf_direct_simple(src), "pdfplumber"


def _fmt_cell(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):  # datetime/date
        return value.isoformat()
    s = str(value).replace("\n", " ").replace("|", "\\|")
    return s


def _convert_xlsx_direct(src: Path) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from zipfile import BadZipFile
    except ImportError as e:
        raise BackendUnavailable("openpyxl not installed") from e
    try:
        wb_data = load_workbook(str(src), data_only=True, read_only=False)
        wb_form = load_workbook(str(src), data_only=False, read_only=False)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as e:
        raise ConversionError("corrupted", f"workbook cannot be parsed: {e}") from e

    if getattr(wb_form, "security", None) and getattr(
        wb_form.security, "workbookPassword", None
    ):
        raise ConversionError(
            "protected",
            "workbook contains protected or encrypted sheets; "
            "ask sender for unprotected copy",
        )

    out: list[str] = []
    warnings: list[tuple[str, int, int]] = []  # (sheet, charts, pivots)

    for ws_data, ws_form in zip(wb_data.worksheets, wb_form.worksheets):
        if ws_form.protection.sheet:
            raise ConversionError(
                "protected",
                f"sheet '{ws_form.title}' is protected; ask sender for unprotected copy",
            )
        out.append(f"## {ws_form.title}\n")

        max_row = ws_data.max_row or 0
        max_col = ws_data.max_column or 0
        if max_row == 0 or max_col == 0:
            out.append("(empty sheet)\n")
            continue

        merged_ranges = list(ws_form.merged_cells.ranges)
        merged_covered: set[tuple[int, int]] = set()
        for mr in merged_ranges:
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    if (r, c) != (mr.min_row, mr.min_col):
                        merged_covered.add((r, c))

        used = []
        for r in range(1, max_row + 1):
            row_vals = []
            any_val = False
            for c in range(1, max_col + 1):
                if (r, c) in merged_covered:
                    row_vals.append("")
                    continue
                v_data = ws_data.cell(row=r, column=c).value
                v_form = ws_form.cell(row=r, column=c).value
                if v_data is not None:
                    text = _fmt_cell(v_data)
                elif isinstance(v_form, str) and v_form.startswith("="):
                    text = v_form
                else:
                    text = ""
                hl = ws_form.cell(row=r, column=c).hyperlink
                if hl and hl.target:
                    text = f"[{text or hl.target}]({hl.target})"
                if text:
                    any_val = True
                row_vals.append(text)
            if any_val:
                used.append(row_vals)
        if not used:
            out.append("(empty sheet)\n")
            continue

        ncols = max(len(r) for r in used)
        while ncols > 0 and all(len(r) < ncols or r[ncols - 1] == "" for r in used):
            ncols -= 1
        if ncols == 0:
            out.append("(empty sheet)\n")
            continue

        used = [r[:ncols] + [""] * (ncols - len(r)) for r in used]
        out.append("| " + " | ".join(used[0]) + " |")
        out.append("| " + " | ".join(["---"] * ncols) + " |")
        for row in used[1:]:
            out.append("| " + " | ".join(row) + " |")
        out.append("")

        n_charts = len(getattr(ws_form, "_charts", []) or [])
        n_pivots = len(getattr(ws_form, "_pivots", []) or [])
        for _ in range(n_charts):
            out.append(
                f'<!-- chart not rendered: sheet="{ws_form.title}" '
                f'range="" type="unknown" -->'
            )
        for _ in range(n_pivots):
            out.append(
                f'<!-- pivot table not rendered: sheet="{ws_form.title}" range="" -->'
            )
        for img in getattr(ws_form, "_images", []) or []:
            anchor = getattr(img, "anchor", "")
            out.append(
                f'<!-- embedded image: sheet="{ws_form.title}" anchor="{anchor}" -->'
            )
        if n_charts or n_pivots:
            warnings.append((ws_form.title, n_charts, n_pivots))

    if warnings:
        total_charts = sum(c for _, c, _ in warnings)
        total_pivots = sum(p for _, _, p in warnings)
        print(
            f"WARN xlsx-partial: {total_charts} chart(s)/{total_pivots} "
            f"pivot(s) not rendered",
            file=sys.stderr,
        )

    return "\n".join(out).rstrip() + "\n"


def _xlsx_has_charts_or_pivots(src: Path) -> bool:
    """Cheap probe: True iff any sheet contains a chart or a pivot table."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return False
    try:
        wb = load_workbook(str(src), data_only=False, read_only=False)
    except Exception:
        return False
    for ws in wb.worksheets:
        if getattr(ws, "_charts", None) or getattr(ws, "_pivots", None):
            return True
    return False


def _try_markitdown_xlsx(src: Path, use_uv: bool) -> tuple[str, str] | None:
    """Level 1 (uvx) -> Level 2 (pip). Returns (body, backend) or None."""
    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, "xlsx"), "markitdown[xlsx]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed(["markitdown[xlsx]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), "markitdown[xlsx]/pip"
        except BackendUnavailable:
            pass
    return None


def _do_xlsx(src: Path, use_uv: bool) -> tuple[str, str]:
    """XLSX dispatch with full fallback chain per spec section 5."""
    if _xlsx_has_charts_or_pivots(src):
        fallback = _try_markitdown_xlsx(src, use_uv)
        if fallback is not None:
            return fallback

    try:
        return _convert_xlsx_direct(src), "openpyxl"
    except ConversionError as direct_err:
        fallback = _try_markitdown_xlsx(src, use_uv)
        if fallback is not None:
            return fallback
        raise direct_err
    except BackendUnavailable:
        fallback = _try_markitdown_xlsx(src, use_uv)
        if fallback is not None:
            return fallback
        raise ConversionError(
            "install-failed", "xlsx requires openpyxl or markitdown[xlsx]"
        )


def _convert_pptx_direct(src: Path) -> str:
    try:
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError, PythonPptxError
        from lxml.etree import XMLSyntaxError
        from zipfile import BadZipFile
    except ImportError as e:
        raise BackendUnavailable("python-pptx not installed") from e
    try:
        p = Presentation(str(src))
    except (
        PackageNotFoundError,
        PythonPptxError,
        XMLSyntaxError,
        BadZipFile,
        KeyError,
        OSError,
    ) as e:
        raise ConversionError("corrupted", f"PPTX cannot be parsed: {e}") from e
    out: list[str] = []
    for i, slide in enumerate(p.slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        out.append(text)
            elif getattr(shape, "shape_type", None) == 13:  # picture
                out.append(f"<!-- embedded image: slide={i} -->")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("> " + notes.replace("\n", "\n> "))
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def _do_pptx(src: Path, use_uv: bool) -> tuple[str, str]:
    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, "pptx"), "markitdown[pptx]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed(["markitdown[pptx]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), "markitdown[pptx]/pip"
        except BackendUnavailable:
            pass
    try:
        return _convert_pptx_direct(src), "python-pptx"
    except BackendUnavailable as e:
        raise ConversionError(
            "install-failed", "pptx requires python-pptx or markitdown"
        ) from e


def _convert_html_direct(src: Path) -> str:
    try:
        from bs4 import BeautifulSoup, NavigableString
    except ImportError as e:
        raise BackendUnavailable("beautifulsoup4 not installed") from e
    html = src.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()

    def walk(node) -> str:
        if isinstance(node, NavigableString):
            return str(node)
        name = node.name
        children = "".join(walk(c) for c in node.children).strip()
        if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(name[1])
            return "\n" + "#" * level + " " + children + "\n\n"
        if name == "p":
            return "\n" + children + "\n"
        if name == "br":
            return "\n"
        if name == "a":
            href = node.get("href", "")
            return f"[{children}]({href})" if href else children
        if name == "li":
            return "- " + children + "\n"
        if name in {"strong", "b"}:
            return f"**{children}**"
        if name in {"em", "i"}:
            return f"*{children}*"
        if name == "code":
            return f"`{children}`"
        if name == "table":
            rows = []
            for tr in node.find_all("tr"):
                cells = [
                    walk(td).strip().replace("|", "\\|").replace("\n", " ")
                    for td in tr.find_all(["th", "td"])
                ]
                rows.append(cells)
            if not rows:
                return ""
            ncols = max(len(r) for r in rows)
            rows = [r + [""] * (ncols - len(r)) for r in rows]
            out = [
                "| " + " | ".join(rows[0]) + " |",
                "| " + " | ".join(["---"] * ncols) + " |",
            ]
            for r in rows[1:]:
                out.append("| " + " | ".join(r) + " |")
            return "\n" + "\n".join(out) + "\n"
        return children

    body = walk(soup.body or soup)
    while "\n\n\n" in body:
        body = body.replace("\n\n\n", "\n\n")
    return body.strip() + "\n"


def _do_html(src: Path, use_uv: bool) -> tuple[str, str]:
    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, "html"), "markitdown[html]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed(["markitdown[html]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), "markitdown[html]/pip"
        except BackendUnavailable:
            pass
    try:
        return _convert_html_direct(src), "bs4"
    except BackendUnavailable as e:
        raise ConversionError("install-failed", "html requires bs4 or markitdown") from e


def _convert_rtf_direct(src: Path) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as e:
        raise BackendUnavailable("striprtf not installed") from e
    raw = src.read_text(encoding="utf-8", errors="replace")
    text = rtf_to_text(raw).strip()
    return text + "\n"


def _do_markitdown_only(src: Path, use_uv: bool, fmt: str) -> tuple[str, str]:
    if fmt == "rtf" and os.environ.get("MTSD_FORCE_PREINSTALLED") != "1":
        try:
            return _convert_rtf_direct(src), "striprtf"
        except BackendUnavailable:
            if _ensure_pip_installed(["striprtf"]):
                try:
                    return _convert_rtf_direct(src), "striprtf/pip"
                except BackendUnavailable:
                    pass
    if use_uv:
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_uvx(p, fmt), f"markitdown[{fmt}]/uvx"
        except BackendUnavailable:
            pass
    if _ensure_pip_installed([f"markitdown[{fmt}]"]):
        try:
            with _with_safe_argv_path(src) as p:
                return _convert_via_markitdown_pip(p), f"markitdown[{fmt}]/pip"
        except BackendUnavailable:
            pass
    raise ConversionError(
        "install-failed",
        f"{fmt} requires PyPI; no network and not in pre-installed bundle",
    )


def _convert_one(
    path: str, *, no_uv: bool, output_path: Path | None = None
) -> ConversionResult:
    src = Path(path).resolve()
    if not src.is_file():
        return _conversion_failure(src, "not-found", str(src))

    size = src.stat().st_size
    if size > MAX_BYTES:
        return _conversion_failure(
            src, "too-large", f"{size // 1_000_000}MB exceeds 50MB limit"
        )

    ext = src.suffix.lower()
    if ext in UNSUPPORTED_DOC:
        return _conversion_failure(
            src,
            "unsupported",
            ".doc legacy format not supported; please open in "
            "Word/LibreOffice and save as .docx",
        )
    if ext not in SUPPORTED_EXTS:
        return _conversion_failure(src, "unsupported", f"format {ext} not in MVP scope")

    # Sources with the same stem share the first output candidate. Keep their
    # compute/write sequence serialized so batch mode cannot race on foo.md.
    with _output_lock_for(src):
        out = output_path or compute_output_path(src)
        if is_up_to_date(src, out):
            return ConversionResult(src=src, exit_code=0, stdout_line=success_line(out))

        started = time.monotonic()
        try:
            body, backend = dispatch(src, no_uv=no_uv)
        except ConversionError as e:
            log_event(
                src,
                backend="none",
                duration_ms=int((time.monotonic() - started) * 1000),
                size_kb=size // 1000,
                exit_code=1,
            )
            return _conversion_failure(src, e.kind, e.message)
        except Exception as e:  # pragma: no cover
            log_event(
                src,
                backend="none",
                duration_ms=int((time.monotonic() - started) * 1000),
                size_kb=size // 1000,
                exit_code=2,
            )
            return _conversion_failure(
                src,
                "internal",
                f"{type(e).__name__}: {e}; please report at "
                "github.com/strigov/mark-all-the-shit-down/issues",
                exit_code=2,
            )

        marker = marker_line(src, backend.split("/", 1)[0])
        write_output(out, marker, body)
        src_mtime = src.stat().st_mtime
        if out.stat().st_mtime < src_mtime:
            os.utime(out, (src_mtime, src_mtime))
        log_event(
            src,
            backend=backend,
            duration_ms=int((time.monotonic() - started) * 1000),
            size_kb=out.stat().st_size // 1000,
            exit_code=0,
        )
        return ConversionResult(src=src, exit_code=0, stdout_line=success_line(out))


def _positive_int(value: str) -> int:
    try:
        parsed = int(value)
    except ValueError as e:
        raise argparse.ArgumentTypeError("must be a positive integer") from e
    if parsed < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return parsed


def _default_jobs(path_count: int) -> int:
    return max(1, min(path_count, os.cpu_count() or 1, 4))


def _precompute_batch_output_paths(paths: list[str]) -> list[Path | None]:
    reserved: set[Path] = set()
    outputs: list[Path | None] = []
    for path in paths:
        src = Path(path).resolve()
        try:
            size = src.stat().st_size if src.is_file() else -1
        except OSError:
            outputs.append(None)
            continue
        ext = src.suffix.lower()
        if (
            size < 0
            or size > MAX_BYTES
            or ext in UNSUPPORTED_DOC
            or ext not in SUPPORTED_EXTS
        ):
            outputs.append(None)
            continue
        out = compute_output_path(src, reserved=reserved)
        reserved.add(out)
        outputs.append(out)
    return outputs


def _emit_results(results: list[ConversionResult], *, include_path: bool) -> None:
    for result in results:
        if result.stdout_line is not None:
            print(result.stdout_line)
        err = result.stderr_line(include_path=include_path)
        if err is not None:
            print(err, file=sys.stderr)


def _combined_exit_code(results: list[ConversionResult]) -> int:
    if any(r.exit_code == 2 for r in results):
        return 2
    if any(r.exit_code != 0 for r in results):
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        prog="convert.py", description="Convert office docs to Markdown."
    )
    p.add_argument("paths", nargs="*", help="absolute path(s) to source document(s)")
    p.add_argument(
        "--version", action="version", version=f"mark-all-the-stuff-down {__version__}"
    )
    p.add_argument(
        "--no-uv",
        action="store_true",
        help="skip uvx, use pip --user / pre-installed only",
    )
    p.add_argument(
        "--jobs",
        type=_positive_int,
        help="parallel conversions for multiple input files (default: up to 4)",
    )
    args = p.parse_args(argv)

    if not args.paths:
        p.print_usage(sys.stderr)
        return 2

    if len(args.paths) == 1:
        results = [_convert_one(args.paths[0], no_uv=args.no_uv)]
        _emit_results(results, include_path=False)
        return results[0].exit_code

    jobs = args.jobs or _default_jobs(len(args.paths))
    output_paths = _precompute_batch_output_paths(args.paths)
    if jobs == 1:
        results = [
            _convert_one(path, no_uv=args.no_uv, output_path=output_path)
            for path, output_path in zip(args.paths, output_paths)
        ]
    else:
        with concurrent.futures.ThreadPoolExecutor(max_workers=jobs) as executor:
            results = list(
                executor.map(
                    lambda item: _convert_one(
                        item[0], no_uv=args.no_uv, output_path=item[1]
                    ),
                    zip(args.paths, output_paths),
                )
            )
    _emit_results(results, include_path=True)
    return _combined_exit_code(results)


if __name__ == "__main__":
    sys.exit(main())
